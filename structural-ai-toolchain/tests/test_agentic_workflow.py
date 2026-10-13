"""The agentic ingestion workflow: atomise → research → clarify inline →
reconstruct → verify → reflect, looping to convergence."""

import pytest

from examples.make_fixtures import make_crack_width_xlsx, make_messy_beam_xlsx
from src.ingestion.agent import ScriptedInteraction, ingest
from src.ingestion.agent.casefile import CaseFile, Fragment


# ---------------------------------------------------------------- casefile
def test_fragment_dependencies_and_casefile():
    case = CaseFile(source_path="x.xlsx", source_format="xlsx")
    case.add(Fragment(fid="F01", kind="input", name="b", value=100))
    case.add(Fragment(fid="F02", kind="input", name="h", value=200))
    case.add(Fragment(fid="F03", kind="derived", name="A", expression="b*h"))
    case.add(Fragment(fid="F04", kind="derived", name="W",
                      expression="A*h/6"))
    assert case.frag("A").depends_on() == {"b", "h"}
    assert case.dependency_errors() == []
    case.frag("b").status = "discarded"
    assert case.dependency_errors()  # A now misses b


# ------------------------------------------------------- end-to-end (tidy)
def test_agentic_ingest_tidy_sheet_converges(tmp_path):
    xlsx = make_crack_width_xlsx(tmp_path)
    res = ingest(xlsx, tmp_path / "repo",
                 answers={"final-result": "w_k"}, render_pdf=False)
    assert res.converged
    assert res.case.final_result == "w_k"
    # replicates the hand-checked EC2 numbers exactly
    import json

    data = json.loads(open(res.paths["json"]).read())
    assert data["results"]["w_k"] == pytest.approx(0.1807, abs=1e-3)
    # audit trail written
    assert (res.case.iterations >= 1)
    assert "journal" in res.paths and "manifest" in res.paths


# ------------------------------------------------- end-to-end (messy sheet)
def test_agentic_ingest_messy_sheet_asks_and_applies(tmp_path):
    xlsx = make_messy_beam_xlsx(tmp_path)
    channel_answers = {
        # magic numbers were surfaced by the parser; presentation Qs:
        "final-result": "Bending_stress_sigma",
        "acceptance-check": "yes",
        "layout": "compact",
    }
    res = ingest(xlsx, tmp_path / "repo", answers=channel_answers,
                 render_pdf=False)
    # decisions recorded in the case file (audit trail)
    qids = {d.qid for d in res.case.decisions}
    assert "final-result" in qids and "layout" in qids
    assert res.case.layout == "compact"
    # sigma is already governed by the sheet's own check (sigma <= fy),
    # so the agent must NOT bolt on a duplicate '<= 1.0' criterion
    checks = [f.expression for f in res.case.kept("check")]
    assert any("Bending_stress_sigma" in c for c in checks)
    assert not any("<= 1.0" in c for c in checks)


# ---------------------------------------------------- defer mode is honest
def test_defer_mode_never_fabricates(tmp_path):
    xlsx = make_messy_beam_xlsx(tmp_path)
    res = ingest(xlsx, tmp_path / "repo", interactive=False,
                 render_pdf=False)
    # deferred questions are surfaced, not silently answered
    assert isinstance(res.deferred, list)
    for d in res.case.decisions:
        assert d.answered_by in ("deferred", "default", "script", "user")


# ------------------------------------------------------- scripted channel
def test_scripted_interaction_matching():
    from src.ingestion.agent.interaction import Question

    chan = ScriptedInteraction({"q1": "42", "layout": "compact"})
    a1, by1 = chan.ask(Question(qid="q1", stage="s", prompt="value?"))
    assert (a1, by1) == ("42", "script")
    a2, by2 = chan.ask(Question(qid="zz", stage="s",
                                prompt="Report layout choice",
                                default="detailed"))
    assert a2 == "compact"      # matched by prompt substring
    a3, by3 = chan.ask(Question(qid="unknown", stage="s", prompt="??",
                                default="d"))
    assert (a3, by3) == ("d", "default")


# ------------------------------------------------------------- resolution
def test_resolution_answers():
    from src.ingestion.agent.stages import _apply_resolution

    case = CaseFile(source_path="x", source_format="x")
    f = case.add(Fragment(fid="U1", kind="derived", name="P_Rd_2",
                          status="unresolved"))
    _apply_resolution(case, f, "91.4 kN")
    assert f.kind == "input" and f.value == 91.4 and f.unit == "kN"
    assert f.status == "kept"

    g = case.add(Fragment(fid="U2", kind="derived", name="x",
                          status="unresolved"))
    _apply_resolution(case, g, "defer")
    assert g.status == "deferred"

    h = case.add(Fragment(fid="U3", kind="derived", name="junk",
                          status="unresolved"))
    _apply_resolution(case, h, "discard")
    assert h.status == "discarded"


# ---------------------------------------------------------- BMP conversion
def test_bmp_to_png_conversion():
    import struct

    from src.ingestion.agent.stages import _bmp_to_png

    # tiny 2x2 24-bit BMP (red, green / blue, white), rows padded to 4B
    width, height = 2, 2
    row = width * 3
    pad = (4 - row % 4) % 4
    pixels = (bytes([0, 0, 255]) + bytes([0, 255, 0]) + b"\x00" * pad +
              bytes([255, 0, 0]) + bytes([255, 255, 255]) + b"\x00" * pad)
    header = (b"BM" + struct.pack("<IHHI", 54 + len(pixels), 0, 0, 54) +
              struct.pack("<IiiHHIIiiII", 40, width, height, 1, 24, 0,
                          len(pixels), 2835, 2835, 0, 0))
    png = _bmp_to_png(header + pixels)
    assert png is not None and png[:8] == b"\x89PNG\r\n\x1a\n"


# ------------------------------------------- reviewer fixes reach the tool
def _slipped_stud_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Stud"
    ws["B1"] = "Headed stud resistance"
    ws["B2"] = "Ultimate strength f_u (MPa)"; ws["C2"] = 450
    ws["B3"] = "Stud diameter d (mm)"; ws["C3"] = 19
    ws["B4"] = "Partial factor gamma_V (-)"; ws["C4"] = 1.25
    ws["B5"] = "Design shear P_Ed (kN)"; ws["C5"] = 57.2
    # computed in N but labelled kN: the classic slip
    ws["B6"] = "Resistance P_Rd (kN)"; ws["C6"] = "=0.8*C2*PI()*C3^2/4/C4"
    ws["B7"] = "Check"; ws["C7"] = "=C5<=C6"
    path = tmp_path / "stud_slip.xlsx"
    wb.save(path)
    return path


def test_reviewer_correction_is_written_into_the_generated_tool(tmp_path):
    import json

    res = ingest(_slipped_stud_xlsx(tmp_path), tmp_path / "repo",
                 answers={"final-result": "Resistance_P_Rd"}, render_pdf=False)
    data = json.loads(open(res.paths["json"]).read())
    assert data["results"]["Resistance_P_Rd"] == pytest.approx(81.66, abs=0.01)
    assert data["review_log"][0]["action"] == "convert"
    code = open(res.paths["python"]).read()
    assert "/1000" in code and "Reviewer correction" in code
    # the generated tool rebuilds the corrected sheet on its own
    ns: dict = {"__file__": res.paths["python"]}
    exec(code, ns)
    rebuilt = ns["build_sheet"]()
    assert rebuilt.results()["Resistance_P_Rd"] == pytest.approx(81.66, abs=0.01)
    assert rebuilt.review_log
    assert res.review is not None and res.review.passed()


# ------------------------------------------------- knowledge-base gate
def test_reconversion_becomes_a_new_revision(tmp_path):
    import json

    xlsx = make_crack_width_xlsx(tmp_path)
    repo = tmp_path / "repo"
    ingest(xlsx, repo, answers={"final-result": "w_k"}, render_pdf=False)
    res = ingest(xlsx, repo, answers={"final-result": "w_k"}, render_pdf=False)
    manifest = json.loads(open(res.paths["manifest"]).read())
    assert manifest["revision"] == "0.2"
    assert manifest["schema"] == "structural-ai-toolchain/manifest@2"


def _renamed_copy(xlsx, tmp_path):
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    wb.active["A1"] = "Crack width check (copy from old project)"
    out = tmp_path / "copy.xlsx"
    wb.save(out)
    return out


def test_duplicate_is_stopped_at_the_gate(tmp_path):
    xlsx = make_crack_width_xlsx(tmp_path)
    repo = tmp_path / "repo"
    ingest(xlsx, repo, answers={"final-result": "w_k"}, render_pdf=False)
    copy = _renamed_copy(xlsx, tmp_path)
    res = ingest(copy, repo, answers={"gatekeeper": "abort"},
                 render_pdf=False)
    assert res.aborted and res.paths == {}
    assert res.case.gatekeeper and res.case.gatekeeper[0]["verdict"] in (
        "duplicate", "similar")
    assert len(list(repo.glob("*/manifest.json"))) == 1


def test_duplicate_can_be_merged_as_a_revision(tmp_path):
    import json

    xlsx = make_crack_width_xlsx(tmp_path)
    repo = tmp_path / "repo"
    first = ingest(xlsx, repo, answers={"final-result": "w_k"},
                   render_pdf=False)
    copy = _renamed_copy(xlsx, tmp_path)
    res = ingest(copy, repo, answers={"gatekeeper": "revise",
                                      "final-result": "w_k"},
                 render_pdf=False)
    assert res.paths["manifest"] == first.paths["manifest"]
    assert json.loads(open(res.paths["manifest"]).read())["revision"] == "0.2"


# ------------------------------------------------------------ AI research
def test_ai_research_assigns_clauses_and_flags_missing_steps(tmp_path):
    from src.llm import ScriptedLLMClient

    xlsx = make_messy_beam_xlsx(tmp_path)
    research = {
        "standard": "EN 1993-1-1",
        "fragments": [{"name": "Bending_stress_sigma",
                       "clause": "EN 1993-1-1 6.2.5 (6.12)",
                       "confidence": "high",
                       "note": "elastic check; class 3 section assumed"},
                      {"name": "Moment_M", "clause": "EN 1993-1-1 6.2.5",
                       "confidence": "high", "note": ""}],
        "missing_steps": [{"description": "lateral-torsional buckling",
                           "clause": "EN 1993-1-1 6.3.2"}],
    }
    review = {"verdict": "APPROVE_WITH_COMMENTS", "findings": [],
              "corrections": []}
    llm = ScriptedLLMClient({"research": research, "review": review})
    res = ingest(xlsx, tmp_path / "repo", llm=llm, render_pdf=False,
                 answers={"final-result": "Bending_stress_sigma",
                          "layout": "detailed"})
    sigma = res.case.frag("Bending_stress_sigma")
    assert sigma.clause == "EN 1993-1-1 6.2.5 (6.12)"
    assert sigma.clause_confidence == "ai-high"
    # a library atom outranks the model; the disagreement is on record
    moment = res.case.frag("Moment_M")
    assert moment.clause_confidence == "atom"
    assert any("AI suggests" in e for e in moment.evidence)
    assert any("lateral-torsional" in r for r in res.case.reflection)
    assert res.case.ai.startswith("scripted")
    assert [c.task for c in llm.calls][:1] == ["research"]
