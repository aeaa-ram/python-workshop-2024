"""Excel parser for The Grinder.

Reads structured calculation workbooks (.xlsx/.xlsm) that follow the
ingestion convention (see docs/ARCHITECTURE.md):

    <Title cell A1>
    INPUTS
    Name | Value | Unit | Description
    ...
    CALCULATIONS
    Name | Formula (real Excel formula) | Unit | Description
    ...
    CHECKS
    Expression (e.g. "w_k <= w_max") | Description

The parser reads the *formulas*, not Excel's cached values: cell
references are mapped back to the variable names declared in the same
sheet, and Excel functions/operators are translated to our expression
syntax. The toolchain then re-computes everything — which doubles as an
independent verification of the legacy workbook.

Free-form spreadsheets that don't follow the convention are parsed
best-effort and produce warnings instead of silent garbage.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from src.ingestion.base import BaseParser
from src.models.calculation import ParsedTool, ParsedVariable

_CELL_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_FUNCTION_MAP = {
    "SQRT": "sqrt",
    "MIN": "min",
    "MAX": "max",
    "ABS": "abs",
    "POWER": "pow",  # POWER(a,b) handled by ** rewrite below
    "LN": "log",
    "EXP": "exp",
    "PI()": "pi",
}
_MARKERS = ("INPUTS", "CALCULATIONS", "CHECKS")


def has_convention_markers(path: Path) -> bool:
    """True if the first sheet uses the strict INPUTS/CALCULATIONS blocks."""
    try:
        wb = load_workbook(path, data_only=False, read_only=True)
        ws = wb.worksheets[0]
        seen = set()
        for row in ws.iter_rows(min_row=1, max_row=200, max_col=1):
            v = row[0].value
            if isinstance(v, str) and v.strip().upper() in _MARKERS:
                seen.add(v.strip().upper())
        wb.close()
        return {"INPUTS", "CALCULATIONS"} <= seen
    except Exception:
        return False


class ConventionExcelParser(BaseParser):
    """Deterministic fast-path for sheets that already follow the strict
    INPUTS/CALCULATIONS/CHECKS convention. Not registered on its own — the
    AIExcelParser delegates here when the markers are present (exact, cheap,
    no LLM). Kept as the reference for well-formed sheets."""

    extensions = ()  # not auto-registered; used via AIExcelParser

    def parse(self, path: Path) -> ParsedTool:
        path = Path(path)
        wb = load_workbook(path, data_only=False)
        ws = wb.worksheets[0]
        tool = ParsedTool(
            title=str(ws["A1"].value or path.stem),
            source_path=str(path),
            source_format=path.suffix.lstrip("."),
        )
        if ws["A2"].value:
            tool.description = str(ws["A2"].value)

        cell_to_name: dict[str, str] = {}
        mode: str | None = None
        for row in ws.iter_rows(min_row=1):
            first = row[0].value
            if isinstance(first, str) and first.strip().upper() in _MARKERS:
                mode = first.strip().upper()
                continue
            if first is None or mode is None:
                continue
            label = str(first).strip()
            if label.lower() == "name" or label.lower() == "expression":
                continue  # header row
            if mode == "INPUTS":
                self._parse_input(tool, row, cell_to_name)
            elif mode == "CALCULATIONS":
                self._parse_calc(tool, row, cell_to_name, ws.title)
            elif mode == "CHECKS":
                self._parse_check(tool, row)
        if not tool.variables:
            tool.warnings.append(
                f"No INPUTS/CALCULATIONS blocks found in {path.name}; the "
                "workbook does not follow the ingestion convention and "
                "needs manual conversion."
            )
        return tool

    # ------------------------------------------------------------------ #
    def _parse_input(self, tool, row, cell_to_name) -> None:
        name = str(row[0].value).strip()
        value_cell = row[1]
        try:
            value = float(value_cell.value)
        except (TypeError, ValueError):
            tool.warnings.append(
                f"Input '{name}': non-numeric value "
                f"{value_cell.value!r} skipped."
            )
            return
        cell_to_name[value_cell.coordinate] = name
        tool.variables.append(
            ParsedVariable(
                name=name,
                value=value,
                unit=str(row[2].value or "") if len(row) > 2 else "",
                description=str(row[3].value or "") if len(row) > 3 else "",
                role="input",
            )
        )

    def _parse_calc(self, tool, row, cell_to_name, sheet_title) -> None:
        name = str(row[0].value).strip()
        formula_cell = row[1]
        raw = formula_cell.value
        if raw is None:
            return
        raw = str(raw)
        if not raw.startswith("="):
            tool.warnings.append(
                f"Calculation '{name}': cell {formula_cell.coordinate} holds "
                f"a literal, not a formula; treated as input."
            )
            try:
                value = float(raw)
            except ValueError:
                return
            cell_to_name[formula_cell.coordinate] = name
            tool.variables.append(
                ParsedVariable(name=name, value=value, role="input")
            )
            return
        expression, unresolved = self.translate_formula(
            raw[1:], cell_to_name
        )
        if unresolved:
            tool.warnings.append(
                f"Calculation '{name}': unresolved cell references "
                f"{unresolved} — formula kept but will not evaluate."
            )
        cell_to_name[formula_cell.coordinate] = name
        tool.variables.append(
            ParsedVariable(
                name=name,
                unit=str(row[2].value or "") if len(row) > 2 else "",
                description=str(row[3].value or "") if len(row) > 3 else "",
                role="derived",
                expression=expression,
            )
        )

    def _parse_check(self, tool, row) -> None:
        tool.checks.append(str(row[0].value).strip())

    @staticmethod
    def translate_formula(
        formula: str, cell_to_name: dict[str, str]
    ) -> tuple[str, list[str]]:
        """Excel formula body -> toolchain expression string."""
        unresolved: list[str] = []

        def repl(match: re.Match) -> str:
            coord = f"{match.group(1)}{match.group(2)}"
            if coord in cell_to_name:
                return cell_to_name[coord]
            unresolved.append(coord)
            return coord

        out = _CELL_REF.sub(repl, formula)
        out = out.replace("^", "**")
        for excel_fn, py_fn in _FUNCTION_MAP.items():
            out = re.sub(
                rf"\b{re.escape(excel_fn)}\b", py_fn, out,
                flags=re.IGNORECASE,
            )
        # POWER(a, b) -> (a)**(b)
        out = re.sub(
            r"\bpow\(([^,]+),([^)]+)\)", r"(\1)**(\2)", out
        )
        return out.strip(), unresolved


# Back-compat alias: tests and other modules import ``ExcelParser`` and use
# its ``translate_formula`` staticmethod.
ExcelParser = ConventionExcelParser


# ====================================================================== #
# AI Grinder: cross-sheet formula translation + dispatching parser
# ====================================================================== #
# Excel functions we can translate to sympy-safe syntax.
_TRANSLATABLE_FUNCS = {
    "SQRT": "sqrt", "MIN": "min", "MAX": "max", "ABS": "abs",
    "LN": "log", "LOG": "log", "EXP": "exp",
    "SIN": "sin", "COS": "cos", "TAN": "tan",
    "COT": "cot", "SEC": "sec", "CSC": "csc",
    "ASIN": "asin", "ACOS": "acos", "ATAN": "atan", "ACOT": "acot",
    "SINH": "sinh", "COSH": "cosh", "TANH": "tanh",
}
# sympy-safe function names allowed to remain after translation (the
# catch-all uses this to spot Excel functions we did NOT handle).
_ALLOWED_MATH_FUNCS = set(_TRANSLATABLE_FUNCS.values()) | {"pi"}
# Functions that need external data / lookups — cannot be translated to a
# closed-form expression; trigger human-in-the-loop clarification.
_LOOKUP_FUNCS = ("INDEX", "VLOOKUP", "HLOOKUP", "MATCH", "OFFSET", "INDIRECT",
                 "LOOKUP", "SUMPRODUCT", "SUMIF", "COUNTIF", "XLOOKUP")
_FUNC_CALL = re.compile(r"\b([A-Za-z_][A-Za-z0-9_]*)\s*\(")
_QUALIFIED_REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_]+))!\$?([A-Z]{1,3})\$?(\d+)")
_BARE_REF = re.compile(r"(?<![A-Za-z0-9_!'])\$?([A-Z]{1,3})\$?(\d+)")


class TranslationResult:
    def __init__(self) -> None:
        self.expression = ""
        self.unresolved_refs: list[str] = []
        self.unknown_funcs: list[str] = []
        self.is_check = False
        self.check_expression = ""

    @property
    def ok(self) -> bool:
        return not self.unresolved_refs and not self.unknown_funcs


def translate_excel_formula(
    formula: str, current_sheet: str, cell_map: dict[str, str]
) -> TranslationResult:
    """Translate an Excel formula to toolchain math using a cell->symbol map.

    Handles cross-sheet refs (``Sheet2!B4``, ``'My Sheet'!B4``) and bare
    refs (resolved against ``current_sheet``). Unknown lookup functions and
    unmapped references are reported, not silently dropped, so the caller
    can raise a precise clarification.
    """
    result = TranslationResult()
    body = formula[1:] if formula.startswith("=") else formula
    # openpyxl prefixes newer functions: =_xlfn.COT(...), =_xlfn._xlws.X(...)
    body = re.sub(r"_xlfn\.(?:_xlws\.)?", "", body)

    # text / display formulas ("A"&B) are not calculations
    if "&" in body or '"' in body:
        result.unknown_funcs.append("TEXT")
        result.expression = ""
        return result

    for fn in _LOOKUP_FUNCS:
        if re.search(rf"\b{fn}\s*\(", body, re.IGNORECASE):
            result.unknown_funcs.append(fn)

    # qualified refs first (Sheet!Coord)
    def repl_qualified(m: re.Match) -> str:
        sheet = m.group(1) or m.group(2)
        coord = f"{m.group(3)}{m.group(4)}"
        ref = f"{sheet}!{coord}"
        if ref in cell_map:
            return cell_map[ref]
        result.unresolved_refs.append(ref)
        return "__UNRESOLVED__"

    body = _QUALIFIED_REF.sub(repl_qualified, body)

    # bare refs -> current sheet
    def repl_bare(m: re.Match) -> str:
        coord = f"{m.group(1)}{m.group(2)}"
        ref = f"{current_sheet}!{coord}"
        if ref in cell_map:
            return cell_map[ref]
        result.unresolved_refs.append(ref)
        return "__UNRESOLVED__"

    body = _BARE_REF.sub(repl_bare, body)

    # a check? (top-level comparison, or IF(cond, ...))
    if_match = re.match(r"\s*IF\s*\((.+)\)\s*$", body, re.IGNORECASE)
    cond = None
    if if_match:
        cond = _split_if_condition(if_match.group(1))
    elif any(op in body for op in ("<=", ">=", "<", ">")):
        cond = body
    if cond and any(op in cond for op in ("<=", ">=", "<", ">")):
        result.is_check = True
        result.check_expression = _clean_expr(cond)

    result.expression = _clean_expr(_apply_funcs(body))
    # catch-all: any Excel function we could not translate remains as FUNC(
    leftover = _unhandled_functions(result.expression)
    for fn in leftover:
        if fn not in result.unknown_funcs:
            result.unknown_funcs.append(fn)
    if result.is_check:
        for fn in _unhandled_functions(_apply_funcs(result.check_expression)):
            if fn not in result.unknown_funcs:
                result.unknown_funcs.append(fn)
        result.check_expression = _clean_expr(_apply_funcs(result.check_expression))
    return result


def _split_if_condition(inner: str) -> str:
    """Extract the condition from an IF(cond, a, b) argument list."""
    depth = 0
    for i, ch in enumerate(inner):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        elif ch == "," and depth == 0:
            return inner[:i]
    return inner


def _apply_funcs(body: str) -> str:
    out = body.replace("^", "**")
    # ROUND(x, n) -> x  (drop rounding for the symbolic form)
    out = re.sub(r"\bROUND(?:UP|DOWN)?\s*\(([^,]+),[^)]*\)", r"(\1)",
                 out, flags=re.IGNORECASE)
    out = re.sub(r"\bPOWER\s*\(([^,]+),([^)]+)\)", r"(\1)**(\2)",
                 out, flags=re.IGNORECASE)
    out = re.sub(r"\bPI\s*\(\s*\)", "pi", out, flags=re.IGNORECASE)
    # rounding/truncation dropped for the symbolic form: F(x[,n]) -> (x)
    out = re.sub(r"\b(?:ROUND(?:UP|DOWN)?|TRUNC|INT)\s*\(([^,()]+)(?:,[^)]*)?\)",
                 r"(\1)", out, flags=re.IGNORECASE)
    for xl, py in _TRANSLATABLE_FUNCS.items():
        out = re.sub(rf"\b{xl}\s*\(", f"{py}(", out, flags=re.IGNORECASE)
    return out


def _unhandled_functions(expression: str) -> list[str]:
    """Any FUNC( left after translation is an Excel function we don't map —
    return their names so the caller can raise a precise clarification
    instead of emitting a broken expression."""
    bad = []
    for m in _FUNC_CALL.finditer(expression):
        name = m.group(1)
        if name.lower() not in _ALLOWED_MATH_FUNCS:
            bad.append(name)
    return bad


def _clean_expr(text: str) -> str:
    return text.strip().strip("=").strip()


class AIExcelParser(BaseParser):
    """The AI Grinder's Excel entry point.

    Dispatch:
    - convention sheets (INPUTS/CALCULATIONS/CHECKS) -> exact deterministic
      path (ConventionExcelParser), preserving quality for tidy sheets;
    - everything else (messy, multi-tab, scattered) -> the AI pipeline:
      extract digest -> interpret (LLM if configured, else heuristic) ->
      deterministic formula translation -> verify -> ParsedTool, with
      human-in-the-loop clarifications for anything ambiguous.
    """

    extensions = (".xlsx", ".xlsm")

    def __init__(self, prefer_client: str | None = None) -> None:
        self.prefer_client = prefer_client

    def parse(self, path: Path) -> ParsedTool:
        path = Path(path)
        if has_convention_markers(path):
            tool = ConventionExcelParser().parse(path)
            tool.interpreter = "convention"
            return tool
        from src.ingestion.ai_pipeline import run_ai_pipeline

        return run_ai_pipeline(path, prefer_client=self.prefer_client)


class CsvParser(BaseParser):
    """CSV variant of the ingestion convention (no cell references).

    Rows: block markers INPUTS/CALCULATIONS/CHECKS, then
    ``name,value_or_expression,unit,description``. Because CSV has no
    formula cells, CALCULATIONS rows hold expression strings that
    reference variable names directly, e.g. ``A,b*h,mm^2,Area``.
    """

    extensions = (".csv",)

    def parse(self, path: Path) -> ParsedTool:
        import csv

        path = Path(path)
        tool = ParsedTool(
            title=path.stem.replace("_", " ").title(),
            source_path=str(path),
            source_format="csv",
        )
        mode: str | None = None
        with open(path, newline="", encoding="utf-8-sig") as fh:
            for row in csv.reader(fh):
                if not row or not row[0].strip():
                    continue
                first = row[0].strip()
                if first.upper() in _MARKERS:
                    mode = first.upper()
                    continue
                if first.lower() in ("name", "expression") or mode is None:
                    continue
                unit = row[2].strip() if len(row) > 2 else ""
                desc = row[3].strip() if len(row) > 3 else ""
                if mode == "INPUTS":
                    try:
                        value = float(row[1])
                    except (IndexError, ValueError):
                        tool.warnings.append(
                            f"Input '{first}': missing/non-numeric value."
                        )
                        continue
                    tool.variables.append(
                        ParsedVariable(
                            name=first, value=value, unit=unit,
                            description=desc, role="input",
                        )
                    )
                elif mode == "CALCULATIONS":
                    if len(row) < 2 or not row[1].strip():
                        continue
                    tool.variables.append(
                        ParsedVariable(
                            name=first, unit=unit, description=desc,
                            role="derived", expression=row[1].strip(),
                        )
                    )
                elif mode == "CHECKS":
                    tool.checks.append(first)
        return tool
