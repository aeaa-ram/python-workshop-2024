"""Generate dummy legacy files for trials and tests.

- ``crack_width_ec2.xlsx``  — EN 1992-1-1 §7.3.4 crack width check, built
  with *real Excel formulas* so the Grinder's formula translation is
  exercised for real.
- ``beam_bending_check.ipynb`` — simple notebook following the line
  conventions (``name = expr  # unit | description``).
- ``messy_beam.xlsx`` — a disorganised sheet for the AI Grinder.
- ``stud_resistance_ec4.mcdx`` — a Mathcad Prime package (EN 1994-1-1
  6.6.3.1 headed studs) using the worksheet XML layout of real Prime
  files: units inside formulas, an ``if`` conditional, ``min``, a display
  unit override, a less-or-equal check and narrative text regions.

Run directly (``python examples/make_fixtures.py``) or via
``python main.py demo``.
"""

from __future__ import annotations

from pathlib import Path


def make_crack_width_xlsx(out_dir: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Crack Width"
    ws["A1"] = "Crack Width Verification - RC Slab (EN 1992-1-1 7.3.4)"
    ws["A2"] = (
        "Characteristic crack width for a reinforced concrete slab under "
        "quasi-permanent loading."
    )

    ws["A4"] = "INPUTS"
    ws.append([])  # keep row math simple: write cells explicitly below
    headers = ["Name", "Value", "Unit", "Description"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=5, column=col, value=header)
    inputs = [
        # (name, value, unit, description) — rows 6..21
        ("b", 1000, "mm", "Section width (per metre strip)"),
        ("h", 300, "mm", "Section depth"),
        ("d", 260, "mm", "Effective depth"),
        ("c_nom", 30, "mm", "Nominal concrete cover"),
        ("phi", 16, "mm", "Bar diameter"),
        ("A_s", 1005, "mm^2", "Tension reinforcement area (5 H16)"),
        ("M_qp", 45, "kNm", "Quasi-permanent bending moment"),
        ("E_s", 200000, "MPa", "Steel modulus of elasticity"),
        ("E_cm", 33000, "MPa", "Concrete secant modulus"),
        ("f_ct_eff", 2.9, "MPa", "Effective concrete tensile strength"),
        ("k_t", 0.4, "-", "Load duration factor (long term)"),
        ("k_1", 0.8, "-", "Bond factor (high-bond bars)"),
        ("k_2", 0.5, "-", "Strain distribution factor (bending)"),
        ("k_3", 3.4, "-", "Code coefficient"),
        ("k_4", 0.425, "-", "Code coefficient"),
        ("w_max", 0.3, "mm", "Allowable crack width (XC exposure)"),
    ]
    for offset, (name, value, unit, desc) in enumerate(inputs):
        row = 6 + offset
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=3, value=unit)
        ws.cell(row=row, column=4, value=desc)

    ws.cell(row=23, column=1, value="CALCULATIONS")
    for col, header in enumerate(headers, start=1):
        ws.cell(row=24, column=col, value=header)
    # Input value cells: b=B6 h=B7 d=B8 c_nom=B9 phi=B10 A_s=B11 M_qp=B12
    # E_s=B13 E_cm=B14 f_ct_eff=B15 k_t=B16 k_1=B17 k_2=B18 k_3=B19
    # k_4=B20 w_max=B21. Calc cells start at B25.
    calcs = [
        ("alpha_e", "=B13/B14", "-", "Modular ratio E_s/E_cm"),
        (
            "x",
            "=(B25*B11/B6)*(SQRT(1+2*B6*B8/(B25*B11))-1)",
            "mm",
            "Neutral axis depth (cracked section)",
        ),
        ("z", "=B8-B26/3", "mm", "Internal lever arm"),
        (
            "sigma_s",
            "=B12*1000000/(B27*B11)",
            "MPa",
            "Steel stress under quasi-permanent loads",
        ),
        (
            "h_c_eff",
            "=MIN(2.5*(B7-B8),(B7-B26)/3,B7/2)",
            "mm",
            "Effective tension zone height",
        ),
        ("rho_p_eff", "=B11/(B6*B29)", "-", "Effective reinforcement ratio"),
        (
            "s_r_max",
            "=B19*B9+B17*B18*B20*B10/B30",
            "mm",
            "Maximum crack spacing",
        ),
        (
            "eps_diff",
            "=MAX((B28-B16*B15/B30*(1+B25*B30))/B13,0.6*B28/B13)",
            "-",
            "Mean strain difference (eps_sm - eps_cm)",
        ),
        ("w_k", "=B31*B32", "mm", "Characteristic crack width"),
    ]
    for offset, (name, formula, unit, desc) in enumerate(calcs):
        row = 25 + offset
        ws.cell(row=row, column=1, value=name)
        ws.cell(row=row, column=2, value=formula)
        ws.cell(row=row, column=3, value=unit)
        ws.cell(row=row, column=4, value=desc)

    ws.cell(row=35, column=1, value="CHECKS")
    ws.cell(row=36, column=1, value="Expression")
    ws.cell(row=36, column=2, value="Description")
    ws.cell(row=37, column=1, value="w_k <= w_max")
    ws.cell(row=37, column=2, value="Crack width within allowable limit")

    out_path = out_dir / "crack_width_ec2.xlsx"
    wb.save(out_path)
    return out_path


def make_beam_notebook(out_dir: Path) -> Path:
    import nbformat

    nb = nbformat.v4.new_notebook()
    nb.cells = [
        nbformat.v4.new_markdown_cell(
            "# Rectangular Beam Bending Check\n"
            "Elastic bending stress verification for a simply supported "
            "rectangular timber beam under uniform load."
        ),
        nbformat.v4.new_code_cell(
            "b = 200  # mm | Section width\n"
            "h = 400  # mm | Section depth\n"
            "L = 5000  # mm | Span\n"
            "q = 4.5  # kN/m | Distributed load (design)\n"
            "f_m = 24  # MPa | Design bending strength"
        ),
        nbformat.v4.new_code_cell(
            "M_max = q*(L/1000)**2/8  # kNm | Maximum bending moment\n"
            "W = b*h**2/6  # mm^3 | Elastic section modulus\n"
            "sigma_m = M_max*1000000/W  # MPa | Maximum bending stress"
        ),
        nbformat.v4.new_code_cell(
            "# CHECKS\n"
            "sigma_m <= f_m"
        ),
    ]
    out_path = out_dir / "beam_bending_check.ipynb"
    nbformat.write(nb, str(out_path))
    return out_path


def make_all(out_dir: str | Path) -> list[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    return [
        make_crack_width_xlsx(out_dir),   # tidy: convention fast-path
        make_beam_notebook(out_dir),
        make_messy_beam_xlsx(out_dir),    # messy: AI Grinder + clarifications
        make_stud_mcdx(out_dir),          # Mathcad Prime, unit-aware
    ]


# ---------------------------------------------------------------- Mathcad
_WS_NS = "http://schemas.mathsoft.com/worksheet50"
_ML_NS = "http://schemas.mathsoft.com/math50"


def _ml_id(name: str, unit: bool = False) -> str:
    return f'<ml:id labels="{"UNIT" if unit else "VARIABLE"}">{name}</ml:id>'


def _ml_real(v) -> str:
    return f"<ml:real>{v}</ml:real>"


def _ml_apply(op: str, *args: str) -> str:
    return f"<ml:apply><ml:{op}/>{''.join(args)}</ml:apply>"


def _ml_call(fname: str, *args: str) -> str:
    return (f'<ml:apply><ml:id labels="FUNCTION">{fname}</ml:id>'
            f"<ml:sequence>{''.join(args)}</ml:sequence></ml:apply>")


def _ml_define(name: str, value: str, display_unit: str = "") -> str:
    if display_unit:
        value = (f"<ml:eval>{value}<ml:unitOverride>"
                 f"{_ml_id(display_unit, unit=True)}</ml:unitOverride></ml:eval>")
    return f"<ml:define>{_ml_id(name)}{value}</ml:define>"


def _qty(v, unit: str) -> str:
    return _ml_apply("mult", _ml_real(v), _ml_id(unit, unit=True))


def make_stud_mcdx(out_dir: Path) -> Path:
    import io
    import zipfile

    v, u = _ml_id, _qty
    texts = {
        "rId1": "Shear resistance of headed studs - EN 1994-1-1 6.6.3.1",
        "rId2": "Introduction: design resistance of a 19 mm headed stud "
                "welded through the deck, governed by stud shank failure "
                "or concrete crushing, checked against the design shear "
                "per stud.",
        "rId3": "Stud diameter",
        "rId4": "Ultimate tensile strength of the stud",
    }
    math = [
        _ml_define("d", u(19, "mm")),
        _ml_define("f_u", u(450, "MPa")),
        _ml_define("h_sc", u(100, "mm")),
        _ml_define("f_ck", u(30, "MPa")),
        _ml_define("E_cm", u(33, "GPa")),
        _ml_define("γ_V", _ml_real(1.25)),
        _ml_define("P_Ed", u(57.2, "kN")),
        # alpha = if(h_sc/d >= 4, 1, 0.2 (h_sc/d + 1))   [eq. 6.20/6.21]
        _ml_define("α", _ml_call(
            "if",
            _ml_apply("greaterOrEqual",
                      _ml_apply("div", v("h_sc"), v("d")), _ml_real(4)),
            _ml_real(1),
            _ml_apply("mult", _ml_real(0.2), _ml_apply(
                "plus", _ml_apply("div", v("h_sc"), v("d")), _ml_real(1))))),
        # P_Rd,1 = 0.8 f_u pi d^2/4 / gamma_V   [eq. 6.18]
        _ml_define("P_Rd_1", _ml_apply("div", _ml_apply(
            "mult", _ml_real(0.8), v("f_u"), '<ml:id labels="CONSTANT">π</ml:id>',
            _ml_apply("div", _ml_apply("pow", v("d"), _ml_real(2)),
                      _ml_real(4))), v("γ_V")), display_unit="kN"),
        # P_Rd,2 = 0.29 alpha d^2 sqrt(f_ck E_cm) / gamma_V   [eq. 6.19]
        _ml_define("P_Rd_2", _ml_apply("div", _ml_apply(
            "mult", _ml_real(0.29), v("α"),
            _ml_apply("pow", v("d"), _ml_real(2)),
            _ml_apply("sqrt", _ml_apply("mult", v("f_ck"), v("E_cm")))),
            v("γ_V")), display_unit="kN"),
        _ml_define("P_Rd", _ml_call("min", v("P_Rd_1"), v("P_Rd_2")),
                   display_unit="kN"),
        # effective height below the head (unit written inside the formula)
        _ml_define("h_eff", _ml_apply("minus", v("h_sc"), u(0.003, "m")),
                   display_unit="mm"),
        _ml_define("UR", _ml_apply("div", v("P_Ed"), v("P_Rd"))),
    ]
    regions = [
        '<region region-id="1" left="0" top="0"><text item-idref="rId1"/></region>',
        '<region region-id="2" left="0" top="20"><text item-idref="rId2"/></region>',
        '<region region-id="3" left="0" top="40"><text item-idref="rId3"/></region>',
    ]
    top = 50
    for i, m in enumerate(math):
        if i == 1:
            regions.append(f'<region region-id="t{i}" left="0" top="{top}">'
                           '<text item-idref="rId4"/></region>')
            top += 5
        regions.append(f'<region region-id="m{i}" left="0" top="{top}">'
                       f"<math>{m}</math></region>")
        top += 10
    # the acceptance check: UR <= 1
    regions.append(
        f'<region region-id="chk" left="0" top="{top}"><math><ml:eval>'
        + _ml_apply("lessOrEqual", v("UR"), _ml_real(1))
        + "</ml:eval></math></region>")
    worksheet = (
        f'<?xml version="1.0" encoding="utf-8"?>'
        f'<worksheet xmlns="{_WS_NS}" xmlns:ml="{_ML_NS}">'
        f"<regions>{''.join(regions)}</regions></worksheet>")

    def xaml_package(text: str) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as inner:
            inner.writestr("Xaml/Document.xaml",
                           f'<Section><Paragraph><Run Text="{text}"/>'
                           "</Paragraph></Section>")
        return buf.getvalue()

    rels = "".join(
        f'<Relationship Target="/xaml/FlowDocument{rid[3:]}.XamlPackage" '
        f'Id="{rid}" Type="urn:text"/>' for rid in texts)
    out_path = out_dir / "stud_resistance_ec4.mcdx"
    with zipfile.ZipFile(out_path, "w") as zf:
        zf.writestr("mathcad/worksheet.xml", worksheet)
        zf.writestr("mathcad/_rels/worksheet.xml.rels",
                    f"<Relationships>{rels}</Relationships>")
        for rid, text in texts.items():
            zf.writestr(f"mathcad/xaml/FlowDocument{rid[3:]}.XamlPackage",
                        xaml_package(text))
    return out_path


def make_messy_beam_xlsx(out_dir: Path) -> Path:
    """A deliberately DISORGANISED sheet: no INPUTS/CALCULATIONS blocks,
    scattered labels, values across two tabs, a lookup, and an unlabelled
    magic number — to exercise the AI Grinder + human-in-the-loop path."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Beam"
    ws["B1"] = "Simply Supported Steel Beam — quick check"
    # scattered label/value pairs (label left of value), not in a block
    ws["B3"] = "Span L (m)"; ws["C3"] = 6.0
    ws["B4"] = "Line load w (kN/m)"; ws["C4"] = 12.0
    ws["E3"] = "Section modulus Wel (cm3)"; ws["F3"] = 324.0
    ws["E4"] = "Yield fy (MPa)"; ws["F4"] = 355.0
    # a formula referencing scattered cells (bending moment, kNm)
    ws["B6"] = "Moment M (kNm)"; ws["C6"] = "=C4*C3^2/8"
    # stress = M/Wel with unit juggling; references another tab's factor
    ws["B7"] = "Bending stress sigma (MPa)"; ws["C7"] = "=C6*1000/F3*Factors!B2"
    # a hardcoded magic number with NO label
    ws["H1"] = 1.0
    # a verification via IF
    ws["B9"] = "Utilisation check"; ws["C9"] = "=IF(C7<=F4,1,0)"
    ws["B10"] = "Capacity check"; ws["C10"] = "=C7<=F4"

    ws2 = wb.create_sheet("Factors")
    ws2["A1"] = "Amplification"; ws2["B2"] = 1.0

    out_path = out_dir / "messy_beam.xlsx"
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    for path in make_all(Path(__file__).parent):
        print(path)
